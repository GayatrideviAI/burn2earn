import tempfile
from pathlib import Path
from unittest import mock

from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from apps.accounts.models import Role, User

from . import i18n_tools as tools
from . import services
from .models import CustomerMember, JobRole
from .tests import StockCase


class LanguagePagesTests(StockCase):
    def setUp(self):
        super().setUp()
        user = User.objects.create_user(username="cust1", password="pw-12345-x", role=Role.BUYER)
        CustomerMember.objects.create(user=user, buyer=self.buyer, job_role=JobRole.OPERATIONS)
        self.client.login(username="cust1", password="pw-12345-x")

    def switch(self, code):
        response = self.client.post(reverse("set_language"), {"language": code, "next": reverse("customers:dashboard")})
        self.assertEqual(response.status_code, 302)
        return self.client.get(reverse("customers:dashboard"))

    def test_english_is_the_default(self):
        page = self.client.get(reverse("customers:dashboard"))
        self.assertContains(page, "Stock overview")
        self.assertContains(page, '<html lang="en">')

    def test_the_choice_is_remembered_for_later_pages(self):
        self.switch("hi")
        self.assertContains(self.client.get(reverse("customers:contracts")), "अनुबंध और आपूर्ति")

    def test_hindi_and_punjabi_dashboards(self):
        self.assertContains(self.switch("hi"), "स्टॉक की स्थिति")
        self.assertContains(self.switch("pa"), "ਸਟਾਕ ਦੀ ਸਥਿਤੀ")
        self.assertContains(self.switch("pa"), '<html lang="pa">')

    def test_units_and_choice_labels_are_translated(self):
        page = self.switch("hi")
        self.assertContains(page, "टन")
        self.assertContains(page, "प्लांट संचालन")  # the job role, a model choice label

    def test_alert_texts_are_translated(self):
        # The batch has never been counted, which raises the 'Not counted yet' alert.
        self.assertContains(self.client.get(reverse("customers:dashboard")), "Not counted yet")
        self.assertContains(self.switch("hi"), "अभी गिनती नहीं हुई")
        self.assertContains(self.switch("pa"), "ਹਾਲੇ ਗਿਣਤੀ ਨਹੀਂ ਹੋਈ")

    def test_a_blank_translation_shows_english(self):
        self.assertContains(self.switch("hi"), "Fuel cover")

    def test_the_language_switch_is_on_every_page(self):
        page = self.client.get(reverse("customers:dashboard"))
        for name in ("English", "हिन्दी", "ਪੰਜਾਬੀ"):
            self.assertContains(page, name)

    def test_services_text_follows_the_active_language(self):
        from django.utils import translation

        with translation.override("hi"):
            titles = [a[1] for a in services.build_actions(services.summary(self.buyer), 18)]
        self.assertIn("अभी गिनती नहीं हुई", titles)


class TranslationFilesTests(SimpleTestCase):
    def test_every_marked_phrase_is_in_both_language_files(self):
        found = tools.collect()
        for lang in tools.LANGS:
            missing = [k for k in found if k not in tools.read_po(tools.po_path(lang))]
            self.assertEqual(missing, [], f"Run 'python manage.py translations extract' ({lang}).")

    def test_compiled_files_match_the_po_files(self):
        for lang in tools.LANGS:
            self.assertEqual(tools.mo_path(lang).read_bytes(), tools.build_mo(lang), f"Run 'python manage.py translations compile' ({lang}).")

    def test_translations_keep_the_placeholders(self):
        import re

        pattern = re.compile(r"%\(\w+\)s")
        for lang in tools.LANGS:
            for english, local in tools.read_po(tools.po_path(lang)).items():
                if local:
                    self.assertEqual(sorted(pattern.findall(english)), sorted(pattern.findall(local)), english)
                    self.assertEqual(english.count("%%"), local.count("%%"), english)

    def test_the_sheet_round_trips_and_rejects_broken_placeholders(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(tools, "LOCALE_DIR", Path(tmp)):
            tools.extract()
            sheet = Path(tmp) / "t.xlsx"
            tools.export_sheet(sheet)
            wb = load_workbook(sheet)
            ws = wb["Translations"]
            for row in ws.iter_rows(min_row=2):
                if row[0].value == "Stock":
                    row[1].value = "स्टॉक"
                if row[0].value == "%(n)s days":
                    row[1].value = "दिन"  # lost the placeholder
            wb.save(sheet)
            problems, unknown = tools.import_sheet(sheet)
            self.assertEqual(len(problems), 1)
            self.assertEqual(unknown, [])
            hi = tools.read_po(tools.po_path("hi"))
            self.assertEqual(hi["Stock"], "स्टॉक")
            self.assertEqual(hi["%(n)s days"], "")
