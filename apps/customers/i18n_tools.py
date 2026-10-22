"""Translation tooling for the customer pages.

Plain Python, no gettext programs needed (they are not installed in the Docker image).
Used through ``python manage.py translations <command>``.

    extract   find every marked phrase in the templates and code, and update locale/<lang>/LC_MESSAGES/django.po
    export    write locale/translations.xlsx: English, Hindi, Punjabi columns for people to check and fill in
    import    read that sheet back into the .po files (empty cells stay empty and show English)
    compile   build the django.mo files that Django actually reads

The .po files are the record of the translations; the .mo files are built from them.
"""
import ast
import json
import re
import struct
from pathlib import Path

LANGS = {
    "hi": "Hindi",
    "pa": "Punjabi",
}
PLURAL_FORMS = {"hi": "nplurals=2; plural=(n > 1);", "pa": "nplurals=2; plural=(n > 1);"}

# Notes for the people checking the sheet, keyed by English phrase. Anything not listed has no note.
NOTES = {
    "Pickup": "A visit where the Mitra collects crop residue from a farmer and pays for it.",
    "Log a pickup": "Record that residue was collected from a farmer.",
    "Residue": "Crop residue: straw and stalks left after harvest (paddy straw, etc.).",
    "GeoCell": "A mapped grid square used to group plots; keep the English term if farmers know it that way.",
    "Plot": "A farmer's field, recorded with its area in acres and GPS location.",
    "Land": "Heading for the farmer's plots.",
    "Dispatch": "Sending finished product (pellets, briquettes) from the Grid to a buyer's factory.",
    "Send product to a buyer": "Starts a dispatch to a buyer's factory.",
    "Report output": "Tell the system how much pellets/briquettes were made.",
    "Made": "Quantity of finished product made (pellets, briquettes).",
    "Tonnes on hand": "Residue collected but not yet used up in making product.",
    "Variety": "The crop variety, for example a paddy variety name.",
    "Register a farmer": "Add a new farmer. Needs their consent.",
    "The farmer has been told how their information will be used and agrees.": "Consent statement the Mitra confirms before registering; wording needs legal review.",
    "Sharecropped": "Farmed by someone who gives part of the crop to the landowner.",
    "Tubewell": "Irrigation from a borewell.",
    "Rain-fed": "No irrigation; depends on rain.",
    "Ledger": "The permanent record of each stock movement.",
    "To send": "Changes saved on the phone that have not reached the server yet (shown as a count).",
    "Waiting to send": "Saved on this phone, will be sent when there is a signal.",
    "waiting to send": "Saved on this phone, will be sent when there is a signal.",
    "Online": "Phone has a signal.",
    "Offline": "Phone has no signal; everything is saved on the phone.",
    "t": "Short for tonne (1,000 kg), written after every weight, e.g. '12.5 t'.",
    "Grids": "A Grid is a local collection and processing point run by a Biomass Mitra.",
    "Biomass Mitra": "The local person who collects crop residue from farmers and runs a Grid. Use the word the field team uses.",
    "Biomass Mitras": "Plural of 'Biomass Mitra'.",
    "Farmer": "Shown as a code (e.g. FRM-000012) plus village; never a name.",
    "Fuel cover": "How many days the plant can run on the stock it has available.",
    "Book inventory": "Stock according to the paperwork/records, before anyone counts it.",
    "Physical inventory": "Stock that has actually been counted, plus what moved since.",
    "Usable inventory": "Counted stock that is dry enough for the plant (below its moisture limit).",
    "Available to use": "Usable stock minus what is already set aside for something.",
    "Reconcile": "Action title: the records and the physical count do not match, someone should check.",
    "Chain of custody": "The trail from farmer to Mitra to Grid to truck to the customer's yard.",
    "Tamper-evident trail": "A log where altering an old entry would be detected.",
    "Pile": "A heap of stored biomass inside a yard.",
    "Open to rain": "Stored without a cover so rain can reach it.",
    "Verified": "Counted recently and measured.",
    "Calculated": "Worked out from records and movements, not counted.",
    "Estimated": "Estimated by survey or drone, not weighed.",
    "Overdue for check": "Last counted too long ago; needs a fresh count.",
    "Bales": "Compressed rectangular blocks of straw.",
    "Loose chopped straw": "Straw cut up but not baled.",
    "Pellets": "Small pressed pieces of biomass.",
    "Briquettes": "Larger pressed blocks of biomass.",
    "Weighbridge": "The truck scale at the plant gate.",
    "Moved to another pile": "Stock moved from one pile to another inside the yard.",
    "Burned at the plant": "Used as fuel.",
    "Spoilage, rain and fire risk": "Who bears the loss if stored stock is spoiled, rained on or burnt.",
    "Penalties": "Contract penalties for late or short delivery.",
    "Purchase orders": "A purchase order is the customer's call for a quantity within a contract.",
    "Delivery late": "Alert title.",
    "Contract behind": "Alert title: less delivered than the time passed would suggest.",
    "Not in a yard yet": "Alert title: delivered to the plant gate but not yet recorded as put into a yard.",
}

ROOT = Path(__file__).resolve().parent.parent.parent
LOCALE_DIR = ROOT / "locale"
SOURCES = [
    ROOT / "apps/customers/templates/customers",
    ROOT / "apps/customers/models.py",
    ROOT / "apps/customers/services.py",
    ROOT / "apps/customers/contracts.py",
    ROOT / "apps/customers/templatetags/customers_i18n.py",
    ROOT / "apps/trade/models.py",
    ROOT / "apps/farms/models.py",
    ROOT / "apps/field/templates/field",
    ROOT / "apps/field/forms.py",
    ROOT / "apps/field/views.py",
    ROOT / "apps/field/api.py",
    ROOT / "apps/field/static/field/app.js",
    ROOT / "apps/yardops/templates/yardops",
    ROOT / "apps/yardops/forms.py",
    ROOT / "apps/yardops/views.py",
    ROOT / "apps/yardops/risk.py",
    ROOT / "apps/yardops/services.py",
    ROOT / "apps/yardops/models.py",
    ROOT / "apps/yardops/api.py",
    ROOT / "apps/yardops/static/yardops/app.js",
]
SKIP = {"trace_report.html", "sw.js", "manifest.webmanifest"}  # staff-only page / not shown to people

TRANS_TAG = re.compile(r"""\{%\s*trans\s+(?P<q>["'])(?P<text>.*?)(?P=q)\s*(?:noop\s*)?%\}""", re.S)
BLOCK_TAG = re.compile(r"\{%\s*blocktrans\b(?P<args>[^%]*)%\}(?P<body>.*?)\{%\s*endblocktrans\s*%\}", re.S)
VAR = re.compile(r"\{\{\s*(\w+)\s*\}\}")
PY_CALL = re.compile(r"""(?<![\w.])(?:_|gt|gettext_lazy|gettext)\(\s*(?P<lit>"(?:[^"\\\n]|\\.)*"|'(?:[^'\\\n]|\\.)*')\s*\)""")


def _trim(text):
    # Same as Django's 'trimmed' option: join lines, drop surrounding whitespace.
    return re.sub(r"\s*\n\s*", " ", text.strip())


def scan_template(path):
    text = path.read_text(encoding="utf-8")
    found = []
    for m in TRANS_TAG.finditer(text):
        found.append((m.group("text"), text.count("\n", 0, m.start()) + 1))
    for m in BLOCK_TAG.finditer(text):
        body = m.group("body")
        if "{% plural %}" in body:
            raise ValueError(f"{path}: plurals are not supported in blocktrans, rewrite the sentence")
        if "trimmed" in m.group("args"):
            body = _trim(body)
        body = VAR.sub(lambda v: "%(" + v.group(1) + ")s", body.replace("%", "%%"))
        found.append((body, text.count("\n", 0, m.start()) + 1))
    return found


def scan_python(path):
    text = path.read_text(encoding="utf-8")
    return [(ast.literal_eval(m.group("lit")), text.count("\n", 0, m.start()) + 1) for m in PY_CALL.finditer(text)]


JS_CALL = re.compile(r"""(?<![\w.$])t\(\s*"((?:[^"\\\n]|\\.)*)\"""")


def scan_js(path):
    text = path.read_text(encoding="utf-8")
    return [(json.loads('"' + m.group(1) + '"'), text.count("\n", 0, m.start()) + 1) for m in JS_CALL.finditer(text)]


def section(place):
    if "apps/field/" in place:
        return "Mitra app"
    if "apps/yardops/" in place:
        return "Yard app"
    if "apps/customers/" in place:
        return "Customer app"
    return "Shared labels"


def sections_of(places):
    found = {section(p) for p in places}
    if len(found) > 1:
        return "Several apps"
    return found.pop()


def collect():
    """Every phrase that needs translating: {english: [places it is used]}."""
    found = {}
    for src in SOURCES:
        files = sorted(src.glob("*.html")) if src.is_dir() else [src]
        for f in files:
            if f.name in SKIP:
                continue
            scan = scan_template if f.suffix == ".html" else scan_js if f.suffix == ".js" else scan_python
            for msgid, line in scan(f):
                if msgid.strip():
                    found.setdefault(msgid, []).append(f"{f.relative_to(ROOT)}:{line}")
    return found


# ---- .po files ---------------------------------------------------------------------------------


def _quote(s):
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n") + '"'


def _unquote(s):
    return ast.literal_eval(s) if s.startswith('"') else s


def po_path(lang):
    return LOCALE_DIR / lang / "LC_MESSAGES" / "django.po"


def mo_path(lang):
    return LOCALE_DIR / lang / "LC_MESSAGES" / "django.mo"


def read_po(path):
    """{english: translation} from a .po file (header and comments ignored)."""
    out, msgid, msgstr, mode = {}, None, None, None
    if not path.exists():
        return out

    def flush():
        if msgid is not None and msgid != "":
            out[msgid] = msgstr or ""

    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line.startswith("msgid "):
            flush()
            msgid, msgstr, mode = _unquote(line[6:]), "", "id"
        elif line.startswith("msgstr "):
            msgstr, mode = _unquote(line[7:]), "str"
        elif line.startswith('"'):
            if mode == "id":
                msgid += _unquote(line)
            elif mode == "str":
                msgstr += _unquote(line)
    flush()
    return out


def write_po(lang, found, translations):
    header = [
        "Project-Id-Version: Regenerix customer app",
        f"Language: {lang}",
        "MIME-Version: 1.0",
        "Content-Type: text/plain; charset=UTF-8",
        "Content-Transfer-Encoding: 8bit",
        f"Plural-Forms: {PLURAL_FORMS[lang]}",
    ]
    lines = [f"# {LANGS[lang]} translations for the customer pages.", "# Edit via the spreadsheet (translations export/import) or directly here.", "", 'msgid ""', 'msgstr ""']
    lines += [_quote(h + "\n") for h in header]
    lines.append("")
    for msgid, places in found.items():
        lines.append("#: " + " ".join(places[:3]))
        if msgid in NOTES:
            lines.append("#. " + NOTES[msgid])
        if "%(" in msgid:
            lines.append("#, python-format")
        lines.append("msgid " + _quote(msgid))
        lines.append("msgstr " + _quote(translations.get(msgid, "")))
        lines.append("")
    path = po_path(lang)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def extract():
    found = collect()
    for lang in LANGS:
        existing = read_po(po_path(lang))
        write_po(lang, found, {k: v for k, v in existing.items() if k in found})
    return found


# ---- .mo files ---------------------------------------------------------------------------------


def build_mo(lang):
    """The binary catalogue Django reads, built from the .po file. Blank translations are left out so English shows."""
    entries = {k: v for k, v in read_po(po_path(lang)).items() if v.strip()}
    header = (
        "Project-Id-Version: Regenerix customer app\n"
        f"Language: {lang}\n"
        "MIME-Version: 1.0\n"
        "Content-Type: text/plain; charset=UTF-8\n"
        "Content-Transfer-Encoding: 8bit\n"
        f"Plural-Forms: {PLURAL_FORMS[lang]}\n"
    )
    entries[""] = header
    keys = sorted(entries, key=lambda k: k.encode("utf-8"))
    ids = b""
    strs = b""
    offsets = []
    for k in keys:
        kb, vb = k.encode("utf-8"), entries[k].encode("utf-8")
        offsets.append((len(ids), len(kb), len(strs), len(vb)))
        ids += kb + b"\0"
        strs += vb + b"\0"
    n = len(keys)
    keystart = 7 * 4 + 16 * n
    valuestart = keystart + len(ids)
    table = b""
    for o1, l1, o2, l2 in offsets:
        table += struct.pack("<ii", l1, keystart + o1)
    for o1, l1, o2, l2 in offsets:
        table += struct.pack("<ii", l2, valuestart + o2)
    head = struct.pack("<Iiiiiii", 0x950412DE, 0, n, 7 * 4, 7 * 4 + 8 * n, 0, 0)
    return head + table + ids + strs


def compile_all():
    for lang in LANGS:
        mo_path(lang).write_bytes(build_mo(lang))


# ---- the spreadsheet ---------------------------------------------------------------------------

SHEET = LOCALE_DIR / "translations.xlsx"
HEADERS = ["English (do not edit)", "Hindi / हिन्दी", "Punjabi / ਪੰਜਾਬੀ", "Status", "App", "What it is / where it appears", "Your comments"]


def export_sheet(path=SHEET):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    found = dict(sorted(collect().items(), key=lambda kv: (sections_of(kv[1]) != "Customer app", sections_of(kv[1]))))
    hi, pa = read_po(po_path("hi")), read_po(po_path("pa"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Translations"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D3A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    blank = PatternFill("solid", fgColor="FFF2CC")
    for msgid, places in found.items():
        h, p = hi.get(msgid, ""), pa.get(msgid, "")
        todo = [not v or v == msgid for v in (h, p)]
        status = "English placeholder: please translate" if any(todo) else "Draft: please check"
        where = NOTES.get(msgid, "")
        if msgid.count("%(") or "%%" in msgid:
            where = (where + " " if where else "") + "Keep the %(name)s parts exactly as they are; they are replaced by numbers or names."
        where += (" " if where else "") + "Used in: " + places[0].split("/")[-1]
        ws.append([msgid, h, p, status, sections_of(places), where, ""])
        row = ws.max_row
        for col, is_todo in ((2, todo[0]), (3, todo[1])):
            if is_todo:
                ws.cell(row, col).fill = blank
    for col, width in zip("ABCDEFG", (55, 45, 45, 22, 14, 60, 30)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    dv = DataValidation(type="list", formula1='"English placeholder: please translate,Draft: please check,Checked in the field,Needs discussion"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:D{ws.max_row}")
    ws2 = wb.create_sheet("How to use")
    for line in (
        "This sheet lists every word and sentence the customer application shows.",
        "Column A is the English. Fill in or correct column B (Hindi) and column C (Punjabi). Yellow cells still show the English word because I was not sure of the local wording; please replace it.",
        "Please do not change column A. Leave a cell empty if you do not want a translation; the page will then show English for that phrase.",
        "Where you see %(something)s, keep it exactly as written: the application puts a number or a name there. You may move it within the sentence.",
        "Use the Status column to mark what you have checked. Use the last column for comments such as 'farmers say X here'.",
        "Weights are shown in tonnes (t). Numbers stay in 0-9 digits.",
        "When you are done, send the sheet back and I will load it into the application.",
    ):
        ws2.append([line])
    ws2.column_dimensions["A"].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(found)


def import_sheet(path=SHEET):
    from openpyxl import load_workbook

    ws = load_workbook(path)["Translations"]
    found = collect()
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            rows[str(row[0])] = (str(row[1] or "").strip(), str(row[2] or "").strip())
    problems = []
    for lang, idx in (("hi", 0), ("pa", 1)):
        merged = {}
        for msgid in found:
            if msgid in rows:
                merged[msgid] = rows[msgid][idx]
                if merged[msgid] and sorted(re.findall(r"%\(\w+\)s", msgid)) != sorted(re.findall(r"%\(\w+\)s", merged[msgid])):
                    problems.append(f"{LANGS[lang]}: placeholders differ for '{msgid[:50]}'")
                    merged[msgid] = ""
        write_po(lang, found, merged)
    unknown = [k for k in rows if k not in found]
    return problems, unknown
